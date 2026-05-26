# 飞行员双盲测试评估分析
# Streamlit + Plotly + python-docx

import streamlit as st
import pandas as pd
import numpy as np
import plotly.express as px
import plotly.graph_objects as go
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

st.set_page_config(page_title="飞行员双盲测试评估分析", layout="wide")

# ------------------ 数据读取 ------------------
@st.cache_data
def load_template(file):
    """读取统一模板Excel，处理重复列名"""
    raw = pd.read_excel(file, header=None)
    
    # 获取前两行作为列名
    h1 = raw.iloc[0].fillna("")
    h2 = raw.iloc[1].fillna("")
    
    # 合并列名
    cols_raw = []
    for a, b in zip(h1, h2):
        if pd.notna(a) and str(a).strip() != "":
            cols_raw.append(str(a).replace("\n", ""))
        else:
            cols_raw.append(str(b).replace("\n", ""))
    
    # 处理重复列名：添加后缀
    cols = []
    col_count = {}
    for c in cols_raw:
        if c == "" or c.isspace():
            c = "unnamed"
        if c in col_count:
            col_count[c] += 1
            cols.append(f"{c}_{col_count[c]}")
        else:
            col_count[c] = 1
            cols.append(c)
    
    # 读取数据（从第3行开始）
    df = raw.iloc[3:].copy()
    df.columns = cols
    
    # 删除全空行
    df = df.dropna(how='all')
    
    # 筛选有姓名的行
    if "姓名" in df.columns:
        df = df[df["姓名"].notna()]
    else:
        name_col = None
        for c in df.columns:
            if "姓名" in c:
                name_col = c
                break
        if name_col:
            df = df[df[name_col].notna()]
    
    return df


def extract_deductions(df):
    """提取所有扣分项并计算总得分"""
    df = df.copy()
    
    # 找出所有扣分列（负数值所在的列）
    deduction_cols = []
    for col in df.columns:
        # 检查该列是否有负数值
        numeric_vals = pd.to_numeric(df[col], errors='coerce')
        if (numeric_vals < 0).any():
            deduction_cols.append(col)
    
    # 计算每条记录的扣分总和
    total_deduction = pd.Series(0, index=df.index)
    deduction_details = {}
    
    for col in deduction_cols:
        numeric_vals = pd.to_numeric(df[col], errors='coerce').fillna(0)
        # 只累加负值
        col_deduction = numeric_vals * (numeric_vals < 0)
        total_deduction += col_deduction
        
        # 记录每条记录的扣分详情
        for idx in df.index:
            if col_deduction[idx] < 0:
                if idx not in deduction_details:
                    deduction_details[idx] = []
                deduction_details[idx].append({
                    '扣分项': col,
                    '扣分值': col_deduction[idx]
                })
    
    df['扣分总和'] = total_deduction
    df['最终得分'] = 100 + df['扣分总和']  # 扣分为负数，所以加负值等于扣分
    
    # 记录扣分项数量
    df['扣分项数量'] = [len(deduction_details.get(idx, [])) for idx in df.index]
    
    # 找出高频扣分项（用于分析）
    all_deductions = []
    for idx, details in deduction_details.items():
        for d in details:
            all_deductions.append(d['扣分项'])
    
    freq_deductions = pd.Series(all_deductions).value_counts().head(20).to_dict()
    
    return df, deduction_cols, freq_deductions


def identify_weak_areas(df, deduction_cols):
    """识别薄弱环节"""
    weak_areas = []
    
    for col in deduction_cols:
        numeric_vals = pd.to_numeric(df[col], errors='coerce')
        # 统计扣分次数和总扣分值
        deduction_count = (numeric_vals < 0).sum()
        total_deduction = numeric_vals[numeric_vals < 0].sum()
        
        if deduction_count > 0:
            weak_areas.append({
                '扣分项': col,
                '出现次数': deduction_count,
                '出现频率': f"{deduction_count/len(df)*100:.1f}%",
                '总扣分值': total_deduction,
                '平均每次扣分': total_deduction / deduction_count if deduction_count > 0 else 0
            })
    
    weak_areas_df = pd.DataFrame(weak_areas).sort_values('出现次数', ascending=False)
    return weak_areas_df


# ------------------ Word报告 ------------------
def build_report(all_data, company_data=None, weak_areas=None):
    """生成Word报告"""
    doc = Document()
    
    # 标题
    title = doc.add_heading('飞行员技能评估分析报告', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"生成时间: {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S')}")
    
    # 整体概况
    doc.add_heading('一、整体概况', level=1)
    doc.add_paragraph(f"评估总人数: {len(all_data)} 人")
    doc.add_paragraph(f"整体平均得分: {all_data['最终得分'].mean():.2f} 分")
    doc.add_paragraph(f"得分中位数: {all_data['最终得分'].median():.2f} 分")
    doc.add_paragraph(f"最高分: {all_data['最终得分'].max():.0f} 分")
    doc.add_paragraph(f"最低分: {all_data['最终得分'].min():.0f} 分")
    doc.add_paragraph(f"得分标准差: {all_data['最终得分'].std():.2f} 分")
    
    # 单位对比
    if company_data is not None and len(company_data) > 0:
        doc.add_heading('二、各单位对比分析', level=1)
        
        # 添加单位统计表格
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        headers = ['单位名称', '人数', '平均分', '最高分', '最低分']
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
        
        for _, row in company_data.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(row['单位名称'])
            cells[1].text = str(row['人数'])
            cells[2].text = f"{row['平均分']:.1f}"
            cells[3].text = f"{row['最高分']:.0f}"
            cells[4].text = f"{row['最低分']:.0f}"
    
    # 技术等级分析
    if '技术等级' in all_data.columns:
        doc.add_heading('三、技术等级分析', level=1)
        tech_stats = all_data.groupby('技术等级')['最终得分'].agg(['count', 'mean', 'min', 'max']).round(2)
        tech_stats = tech_stats.reset_index()
        tech_stats.columns = ['技术等级', '人数', '平均分', '最低分', '最高分']
        
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(['技术等级', '人数', '平均分', '最低分', '最高分']):
            hdr_cells[i].text = h
        
        for _, row in tech_stats.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(row['技术等级'])
            cells[1].text = str(row['人数'])
            cells[2].text = f"{row['平均分']:.1f}"
            cells[3].text = f"{row['最低分']:.0f}"
            cells[4].text = f"{row['最高分']:.0f}"
    
    # 薄弱环节分析
    if weak_areas is not None and len(weak_areas) > 0:
        doc.add_heading('四、共性薄弱环节分析', level=1)
        doc.add_paragraph('以下为扣分出现频率最高的项目，建议重点加强训练：')
        
        for _, row in weak_areas.head(10).iterrows():
            doc.add_paragraph(
                f"• {row['扣分项']}: 出现{row['出现次数']}次 ({row['出现频率']})，"
                f"总扣分{abs(row['总扣分值']):.0f}分"
            )
    
    # 个人成绩
    doc.add_heading('五、个人成绩', level=1)
    
    # 按单位分组显示个人成绩
    if '所属单位' in all_data.columns:
        for company in all_data['所属单位'].unique():
            doc.add_heading(f'{company}', level=2)
            company_data_sub = all_data[all_data['所属单位'] == company]
            personal_data = company_data_sub[['姓名', '技术等级', '最终得分', '扣分项数量']].sort_values('最终得分', ascending=False)
            
            table = doc.add_table(rows=1, cols=4)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            for i, h in enumerate(['姓名', '技术等级', '得分', '扣分项数']):
                hdr_cells[i].text = h
            
            for _, row in personal_data.iterrows():
                cells = table.add_row().cells
                cells[0].text = str(row['姓名'])
                cells[1].text = str(row['技术等级']) if pd.notna(row['技术等级']) else '-'
                cells[2].text = f"{row['最终得分']:.1f}"
                cells[3].text = str(row['扣分项数量'])
    
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


# ------------------ 可视化函数 ------------------
def create_score_distribution(df, title="得分分布"):
    """得分分布图"""
    fig = px.histogram(
        df, x="最终得分", nbins=20, 
        title=title,
        labels={"最终得分": "得分", "count": "人数"},
        color_discrete_sequence=["#2ecc71"]
    )
    fig.update_layout(font_family="Microsoft YaHei, SimHei", height=400)
    return fig


def create_company_comparison(company_stats):
    """公司对比图"""
    fig = go.Figure()
    
    fig.add_trace(go.Bar(
        x=company_stats['单位名称'],
        y=company_stats['平均分'],
        name='平均分',
        marker_color='#3498db',
        text=company_stats['平均分'].round(1),
        textposition='auto'
    ))
    
    fig.update_layout(
        title="各单位平均得分对比",
        xaxis_title="单位",
        yaxis_title="平均得分",
        font_family="Microsoft YaHei, SimHei",
        height=400
    )
    return fig


def create_company_boxplot(df):
    """公司得分箱线图"""
    if '所属单位' not in df.columns:
        return None
    
    fig = px.box(
        df, x="所属单位", y="最终得分",
        title="各单位得分分布对比",
        points="all",
        color_discrete_sequence=["#e67e22"]
    )
    fig.update_layout(font_family="Microsoft YaHei, SimHei", height=450)
    return fig


def create_tech_level_comparison(df):
    """技术等级对比图"""
    if '技术等级' not in df.columns:
        return None
    
    tech_stats = df.groupby('技术等级')['最终得分'].agg(['mean', 'count']).reset_index()
    tech_stats.columns = ['技术等级', '平均分', '人数']
    
    fig = go.Figure()
    fig.add_trace(go.Bar(
        x=tech_stats['技术等级'],
        y=tech_stats['平均分'],
        name='平均分',
        marker_color='#2ecc71',
        text=tech_stats['平均分'].round(1),
        textposition='auto'
    ))
    
    fig.update_layout(
        title="各技术等级平均得分对比",
        xaxis_title="技术等级",
        yaxis_title="平均得分",
        font_family="Microsoft YaHei, SimHei",
        height=400
    )
    return fig


def create_tech_boxplot(df):
    """技术等级箱线图"""
    if '技术等级' not in df.columns:
        return None
    
    fig = px.box(
        df, x="技术等级", y="最终得分",
        title="各技术等级得分分布对比",
        color_discrete_sequence=["#9b59b6"]
    )
    fig.update_layout(font_family="Microsoft YaHei, SimHei", height=450)
    return fig


def create_weak_areas_chart(weak_areas):
    """薄弱环节图表"""
    if weak_areas.empty:
        return None
    
    top10 = weak_areas.head(10)
    
    fig = go.Figure()
    fig.add_trace(go.Bar(
        y=top10['扣分项'],
        x=top10['出现次数'],
        orientation='h',
        marker_color='#e74c3c',
        text=top10['出现次数'],
        textposition='outside'
    ))
    
    fig.update_layout(
        title="高频扣分项TOP10",
        xaxis_title="出现次数",
        yaxis_title="扣分项",
        font_family="Microsoft YaHei, SimHei",
        height=500,
        margin=dict(l=250, r=20, t=50, b=20)
    )
    return fig


def create_score_vs_deductions(df):
    """得分与扣分项数量关系图"""
    fig = px.scatter(
        df, x="扣分项数量", y="最终得分",
        title="得分与扣分项数量关系",
        labels={"扣分项数量": "扣分项数量", "最终得分": "最终得分"},
        color="最终得分",
        color_continuous_scale="RdYlGn",
        hover_data=['姓名', '所属单位', '技术等级'] if '姓名' in df.columns else None
    )
    fig.update_layout(font_family="Microsoft YaHei, SimHei", height=450)
    return fig


def create_company_radar(company_stats):
    """公司雷达图"""
    if company_stats.empty:
        return None
    
    categories = company_stats['单位名称'].tolist()
    values = company_stats['平均分'].tolist()
    
    fig = go.Figure()
    fig.add_trace(go.Scatterpolar(
        r=values,
        theta=categories,
        fill='toself',
        name='平均分',
        line_color='#2ecc71'
    ))
    
    fig.update_layout(
        polar=dict(radialaxis=dict(visible=True, range=[0, 100])),
        title="各单位平均分雷达图",
        font_family="Microsoft YaHei, SimHei",
        height=450
    )
    return fig


# ------------------ 页面 ------------------
st.title("✈️ 飞行员技能评估分析平台")
st.markdown("基于《华东飞行员双盲测试数据采集表》统一模板，客观扣分计算")

# 侧边栏
with st.sidebar:
    st.markdown("### 📂 数据上传")
    uploaded_files = st.file_uploader(
        "上传一个或多个Excel文件（每个文件代表一个公司）", 
        type=["xlsx"], 
        accept_multiple_files=True
    )
    
    st.markdown("---")
    st.markdown("### 📌 评分规则")
    st.markdown("""
    - 基础分: **100分**
    - 每出现一次扣分项，扣除相应分数
    - 最终得分 = 100 + 扣分总和（扣分为负值）
    - 分数越高代表表现越好
    """)
    
    st.markdown("---")
    st.markdown("### 🔍 筛选条件")
    
    # 全局筛选器占位
    global_company_filter = st.selectbox("单位筛选", ["全部"], key="global_company")
    global_tech_filter = st.selectbox("技术等级筛选", ["全部"], key="global_tech")

if uploaded_files:
    all_data_list = []
    company_names = []
    
    progress_bar = st.progress(0)
    status_text = st.empty()
    
    for i, file in enumerate(uploaded_files):
        status_text.text(f"正在处理: {file.name}")
        df = load_template(file)
        df, deduction_cols, freq_deductions = extract_deductions(df)
        
        # 添加公司标识
        company_name = file.name.replace('.xlsx', '').replace('_', ' ')
        df['所属单位'] = company_name
        company_names.append(company_name)
        
        all_data_list.append(df)
        progress_bar.progress((i + 1) / len(uploaded_files))
    
    status_text.text("数据处理完成！")
    
    # 合并所有数据
    all_data = pd.concat(all_data_list, ignore_index=True)
    
    # 识别薄弱环节
    all_deduction_cols = []
    for col in all_data.columns:
        numeric_vals = pd.to_numeric(all_data[col], errors='coerce')
        if (numeric_vals < 0).any():
            all_deduction_cols.append(col)
    
    weak_areas = identify_weak_areas(all_data, all_deduction_cols)
    
    # 更新侧边栏筛选器选项
    companies = ["全部"] + sorted(all_data['所属单位'].unique().tolist())
    tech_levels = ["全部"] + sorted(all_data['技术等级'].dropna().unique().tolist()) if '技术等级' in all_data.columns else ["全部"]
    
    # 使用 session_state 或重新创建筛选器
    with st.sidebar:
        selected_company = st.selectbox("单位筛选", companies, key="company_filter")
        selected_tech = st.selectbox("技术等级筛选", tech_levels, key="tech_filter")
    
    # 应用筛选
    filtered_data = all_data.copy()
    if selected_company != "全部":
        filtered_data = filtered_data[filtered_data['所属单位'] == selected_company]
    if selected_tech != "全部":
        filtered_data = filtered_data[filtered_data['技术等级'] == selected_tech]
    
    # 计算公司统计
    company_stats = all_data.groupby('所属单位')['最终得分'].agg(
        人数='count', 平均分='mean', 最高分='max', 最低分='min', 标准差='std'
    ).round(2).reset_index()
    company_stats.columns = ['单位名称', '人数', '平均分', '最高分', '最低分', '标准差']
    
    # 主界面
    st.success(f"✅ 成功加载 {len(all_data)} 条飞行员记录，来自 {len(company_names)} 个单位")
    
    tab1, tab2, tab3, tab4, tab5 = st.tabs([
        "📊 总体分析", "🏢 单位对比", "📈 技术等级分析", "👨‍✈️ 个人详情", "📄 报告输出"
    ])
    
    with tab1:
        col1, col2, col3, col4 = st.columns(4)
        col1.metric("总评估人数", len(filtered_data))
        col2.metric("整体平均分", f"{filtered_data['最终得分'].mean():.1f}")
        col3.metric("最高分", f"{filtered_data['最终得分'].max():.0f}")
        col4.metric("最低分", f"{filtered_data['最终得分'].min():.0f}")
        
        # 得分分布
        fig1 = create_score_distribution(filtered_data, "得分分布直方图")
        st.plotly_chart(fig1, use_container_width=True)
        
        # 得分与扣分项关系
        fig2 = create_score_vs_deductions(filtered_data)
        st.plotly_chart(fig2, use_container_width=True)
        
        # 统计摘要
        st.markdown("### 统计摘要")
        stats_df = pd.DataFrame({
            "指标": ["平均分", "中位数", "标准差", "最高分", "最低分", "总扣分项数"],
            "数值": [
                f"{filtered_data['最终得分'].mean():.2f}",
                f"{filtered_data['最终得分'].median():.2f}",
                f"{filtered_data['最终得分'].std():.2f}",
                f"{filtered_data['最终得分'].max():.0f}",
                f"{filtered_data['最终得分'].min():.0f}",
                f"{filtered_data['扣分项数量'].sum()}"
            ]
        })
        st.dataframe(stats_df, use_container_width=True)
        
        # 薄弱环节
        st.markdown("### ⚠️ 共性薄弱环节分析")
        st.markdown("以下为扣分出现频率最高的项目，建议重点加强训练：")
        
        if not weak_areas.empty:
            fig3 = create_weak_areas_chart(weak_areas)
            if fig3:
                st.plotly_chart(fig3, use_container_width=True)
            
            st.dataframe(weak_areas.head(15), use_container_width=True, hide_index=True)
        else:
            st.info("未发现扣分项，所有飞行员表现良好！")
    
    with tab2:
        st.markdown("### 各单位对比分析")
        
        # 公司统计表
        st.dataframe(company_stats, use_container_width=True, hide_index=True)
        
        col1, col2 = st.columns(2)
        with col1:
            fig_comp = create_company_comparison(company_stats)
            st.plotly_chart(fig_comp, use_container_width=True)
        
        with col2:
            fig_radar = create_company_radar(company_stats)
            if fig_radar:
                st.plotly_chart(fig_radar, use_container_width=True)
        
        # 箱线图
        fig_box = create_company_boxplot(all_data)
        if fig_box:
            st.plotly_chart(fig_box, use_container_width=True)
        
        # 单位详细分析
        st.markdown("### 各单位详细数据")
        for company in company_stats['单位名称']:
            with st.expander(f"📊 {company} 详细分析"):
                company_data = all_data[all_data['所属单位'] == company]
                
                c1, c2, c3, c4 = st.columns(4)
                c1.metric("人数", len(company_data))
                c2.metric("平均分", f"{company_data['最终得分'].mean():.1f}")
                c3.metric("最高分", f"{company_data['最终得分'].max():.0f}")
                c4.metric("最低分", f"{company_data['最终得分'].min():.0f}")
                
                # 该单位的得分分布
                fig_comp_dist = create_score_distribution(company_data, f"{company} 得分分布")
                st.plotly_chart(fig_comp_dist, use_container_width=True)
                
                # 该单位的高频扣分项
                comp_deduction_cols = []
                for col in company_data.columns:
                    numeric_vals = pd.to_numeric(company_data[col], errors='coerce')
                    if (numeric_vals < 0).any():
                        comp_deduction_cols.append(col)
                
                comp_weak = identify_weak_areas(company_data, comp_deduction_cols)
                if not comp_weak.empty:
                    st.markdown("**高频扣分项TOP5**")
                    st.dataframe(comp_weak.head(5), use_container_width=True, hide_index=True)
    
    with tab3:
        if '技术等级' in all_data.columns:
            st.markdown("### 技术等级分析")
            
            col1, col2 = st.columns(2)
            with col1:
                fig_tech = create_tech_level_comparison(filtered_data)
                if fig_tech:
                    st.plotly_chart(fig_tech, use_container_width=True)
            
            with col2:
                fig_tech_box = create_tech_boxplot(filtered_data)
                if fig_tech_box:
                    st.plotly_chart(fig_tech_box, use_container_width=True)
            
            # 技术等级详细统计
            tech_stats_detail = filtered_data.groupby('技术等级').agg({
                '最终得分': ['count', 'mean', 'std', 'min', 'max'],
                '扣分项数量': 'sum'
            }).round(2)
            tech_stats_detail.columns = ['人数', '平均分', '标准差', '最低分', '最高分', '总扣分次数']
            st.dataframe(tech_stats_detail, use_container_width=True)
            
            # 各技术等级与单位交叉分析
            st.markdown("### 各单位技术等级分布")
            pivot_table = pd.crosstab(
                filtered_data['所属单位'], 
                filtered_data['技术等级'], 
                values=filtered_data['最终得分'], 
                aggfunc='mean'
            ).round(1)
            st.dataframe(pivot_table, use_container_width=True)
        else:
            st.info("数据中无技术等级信息")
    
    with tab4:
        st.markdown("### 个人详情查询")
        
        # 多重筛选
        col1, col2 = st.columns(2)
        with col1:
            company_filter = st.selectbox("选择单位", ["全部"] + sorted(all_data['所属单位'].unique().tolist()), key="personal_company")
        with col2:
            if '技术等级' in all_data.columns:
                tech_filter = st.selectbox("选择技术等级", ["全部"] + sorted(all_data['技术等级'].dropna().unique().tolist()), key="personal_tech")
            else:
                tech_filter = "全部"
        
        # 搜索
        search_name = st.text_input("🔍 搜索姓名", "")
        
        # 应用筛选
        personal_data = all_data.copy()
        if company_filter != "全部":
            personal_data = personal_data[personal_data['所属单位'] == company_filter]
        if tech_filter != "全部":
            personal_data = personal_data[personal_data['技术等级'] == tech_filter]
        if search_name:
            personal_data = personal_data[personal_data['姓名'].astype(str).str.contains(search_name, na=False)]
        
        # 显示表格
        display_cols = ['姓名', '所属单位', '技术等级', '最终得分', '扣分项数量'] if '技术等级' in all_data.columns else ['姓名', '所属单位', '最终得分', '扣分项数量']
        display_cols = [c for c in display_cols if c in personal_data.columns]
        
        st.dataframe(personal_data[display_cols].sort_values('最终得分', ascending=False), 
                     use_container_width=True, hide_index=True)
        
        # 个人详细分析
        if len(personal_data) > 0:
            st.markdown("### 📋 个人详细分析")
            selected_pilot = st.selectbox("选择飞行员", personal_data['姓名'].tolist())
            pilot_data = personal_data[personal_data['姓名'] == selected_pilot].iloc[0]
            
            c1, c2, c3, c4 = st.columns(4)
            c1.metric("最终得分", f"{pilot_data['最终得分']:.1f}")
            c2.metric("扣分项数量", pilot_data['扣分项数量'])
            c3.metric("所属单位", pilot_data['所属单位'] if pd.notna(pilot_data['所属单位']) else "-")
            if '技术等级' in pilot_data.index:
                c4.metric("技术等级", pilot_data['技术等级'] if pd.notna(pilot_data['技术等级']) else "-")
            
            # 显示该飞行员的具体扣分项
            st.markdown("#### 扣分详情")
            pilot_deductions = []
            for col in all_deduction_cols:
                val = pd.to_numeric(pilot_data.get(col, 0), errors='coerce')
                if val < 0:
                    pilot_deductions.append({'扣分项': col, '扣分值': val})
            
            if pilot_deductions:
                deductions_df = pd.DataFrame(pilot_deductions)
                st.dataframe(deductions_df, use_container_width=True, hide_index=True)
            else:
                st.success("该飞行员无扣分项，表现完美！")
            
            # 与同单位平均分对比
            company_avg = all_data[all_data['所属单位'] == pilot_data['所属单位']]['最终得分'].mean()
            st.metric("与单位平均分对比", 
                     f"{pilot_data['最终得分'] - company_avg:+.1f} 分",
                     delta=f"单位平均: {company_avg:.1f}分")
    
    with tab5:
        st.markdown("### 📄 生成分析报告")
        
        # 报告预览
        report_preview = f"""
        ═══════════════════════════════════════════════════════════
                       飞行员技能评估分析报告
        ═══════════════════════════════════════════════════════════
        
        生成时间: {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S')}
        
        【整体概况】
        评估人数: {len(all_data)} 人
        整体平均得分: {all_data['最终得分'].mean():.2f} 分
        得分中位数: {all_data['最终得分'].median():.2f} 分
        最高分: {all_data['最终得分'].max():.0f} 分
        最低分: {all_data['最终得分'].min():.0f} 分
        得分标准差: {all_data['最终得分'].std():.2f} 分
        
        【各单位概况】
        """
        for _, row in company_stats.iterrows():
            report_preview += f"\n  {row['单位名称']}: {row['人数']}人, 平均分{row['平均分']:.1f}分"
        
        if not weak_areas.empty:
            report_preview += f"\n\n【高频扣分项TOP5】"
            for _, row in weak_areas.head(5).iterrows():
                report_preview += f"\n  • {row['扣分项']}: 出现{row['出现次数']}次 ({row['出现频率']})"
        
        report_preview += "\n\n═══════════════════════════════════════════════════════════"
        
        st.text_area("报告预览", report_preview, height=400)
        
        col1, col2 = st.columns(2)
        with col1:
            # 生成Word报告
            report_bio = build_report(all_data, company_stats, weak_areas)
            st.download_button(
                "📥 下载Word报告",
                report_bio,
                f"飞行员评估报告_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        
        with col2:
            # 下载CSV数据
            csv_data = all_data.to_csv(index=False, encoding='utf-8-sig')
            st.download_button(
                "📥 下载CSV数据",
                csv_data,
                f"飞行员评估数据_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}.csv",
                mime="text/csv"
            )
        
        # 导出各公司单独报告
        st.markdown("#### 导出各单位单独报告")
        for company in company_stats['单位名称']:
            company_data = all_data[all_data['所属单位'] == company]
            company_weak = identify_weak_areas(company_data, all_deduction_cols)
            company_report = build_report(company_data, None, company_weak)
            st.download_button(
                f"📥 下载 {company} 报告",
                company_report,
                f"{company}_评估报告.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

else:
    st.info("👈 请先从左侧上传一个或多个Excel文件（每个文件代表一个公司）")
    
    with st.expander("📖 文件格式说明"):
        st.markdown("""
        **支持的文件格式**：
        - 按《华东飞行员双盲测试数据采集表》模板填写的Excel文件
        - 每个文件代表一个公司的数据
        - 文件应包含：序号、姓名、技术等级等基本信息列
        
        **系统会自动**：
        1. 识别所有扣分项（负数值）
        2. 计算每位飞行员的最终得分（100 + 扣分总和）
        3. 统计高频扣分项，识别薄弱环节
        4. 支持多公司对比分析
        5. 可按技术等级、单位等多维度筛选
        
        **评分标准**：
        - 基础分: 100分
        - 每出现一次扣分项，扣除相应分数
        - 最终得分 = 100 + 扣分总和
        """)

st.markdown("---")
st.caption("飞行员双盲测试评估分析 | 基于华东局B737飞行技能评估标准 | 客观扣分计算")